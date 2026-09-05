from __future__ import annotations

from collections.abc import Callable
from pathlib import Path
from copy import copy
import re

from config import DEFAULT_MAPPING, load_mapping_rules
from extraction.edb360_assessment import build_edb360_assessment_mapping
from extraction.edb360_metadata import build_edb360_text_mapping, extract_edb360_metadata
from extraction.extract_html import extract_content_from_input
from mapping.content_registry import ContentRegistry
from mapping.mapper import resolve_mappings
from models import ExtractedContent, GenerationReport, ImageContent, MappingRule, MultiImageContent, OperationCancelled, TableContent
from placeholder_inserter import PlaceholderInsertReport, insert_mapping_placeholders
from rpwithchart import render_excel_report
from rendering.word_renderer import document_contains_placeholder, render_report
from sql_healthcheck.merge_sql import merge_sql_root_healthcheck


LogCallback = Callable[[str], None]
DEFAULT_SQL_MAPPING = Path(__file__).resolve().parent / "mapping" / "sql_healthcheck_mapping.yaml"
DEFAULT_EDB360_MASTER_TEMPLATE = Path(__file__).resolve().parent / "templates" / "edb360_master.docx"
DEFAULT_SQL_MASTER_TEMPLATE = Path(__file__).resolve().parent / "templates" / "sql_healthcheck_master.docx"
DEFAULT_EDB360_MASTER_TEMPLATE_EN = Path(__file__).resolve().parent / "templates" / "edb360_master_en.docx"
DEFAULT_SQL_MASTER_TEMPLATE_EN = Path(__file__).resolve().parent / "templates" / "sql_healthcheck_master_en.docx"
DEFAULT_MAX_TABLE_DATA_ROWS = 50
EDB360_CHART_GROUPS = {
    "<log_switch_charts>": ("log_switch_frequency_for_instance", "Instance {instance}: Log switch frequency"),
    "<aas_per_wait_class_charts>": ("aas_per_wait_class_for_instance", "Instance {instance}: AAS per Wait Class"),
    "<ash_top_timed_events_charts>": ("ash_top_timed_events_for_instance", "Instance {instance}: ASH Top Timed Events"),
    "<ash_top_sql_charts>": ("ash_top_sql_for_cluster", "ASH Top SQL for Cluster"),
    "<cpu_busy_idle_charts>": ("cpu_busy_and_idle_times_percent_for_instance", "Instance {instance}: CPU Busy and Idle Times Percent"),
    "<memory_statistics_charts>": ("memory_statistics_for_instance", "Instance {instance}: Memory Statistics"),
    "<sga_statistics_charts>": ("sga_statistics_for_instance", "Instance {instance}: SGA Statistics"),
    "<pga_statistics_charts>": ("pga_statistics_for_instance", "Instance {instance}: PGA Statistics"),
}


def generate_report(
    html_root_folder: str,
    word_file: str,
    output_file_path: str,
    log_callback: LogCallback | None = None,
    cancel_check: Callable[[], bool] | None = None,
) -> str:
    """Generate a report using the GUI Save As workflow."""
    return generate_report_to_file(
        html_input=_validate_html_root_folder(html_root_folder),
        word_file=word_file,
        output_file=output_file_path,
        mapping_file=DEFAULT_MAPPING,
        chart_output_dir=Path(output_file_path).parent / "generated_charts",
        log_callback=log_callback,
        cancel_check=cancel_check,
    )


def generate_report_to_file(
    html_input: str | Path,
    word_file: str | Path,
    output_file: str | Path,
    mapping_file: str | Path = DEFAULT_MAPPING,
    chart_output_dir: str | Path | None = None,
    text_mapping: dict[str, str] | None = None,
    edb360_one_click: bool = False,
    allow_chart_fallback: bool = True,
    validate_only: bool = False,
    log_callback: LogCallback | None = None,
    cancel_check: Callable[[], bool] | None = None,
) -> str:
    """Generate a report to an explicit output file path."""
    input_path = _validate_html_input(html_input)
    template_path = _validate_word_file(word_file)
    output_path = Path(output_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if output_path.suffix.lower() != ".docx":
        raise ValueError(f"Output file must be a .docx file: {output_path}")

    mapping_path = Path(mapping_file)
    if not mapping_path.is_file():
        raise FileNotFoundError(f"Mapping file does not exist: {mapping_path}")

    chart_dir = Path(chart_output_dir) if chart_output_dir else output_path.parent / "generated_charts"
    report = GenerationReport()

    log = _make_logger(log_callback)
    log("Validating inputs...")
    log(f"HTML input: {input_path}")
    log(f"Template: {template_path}")
    log(f"Mapping: {mapping_path}")
    log(f"Output: {output_path}")

    log("Loading mapping rules...")
    rules = load_mapping_rules(mapping_path)
    log(f"Loaded mapping rules: {len(rules)}")
    _check_cancelled(cancel_check)

    chart_rules = _chart_rules_present_in_template(template_path, rules)
    log(f"Chart placeholders in template: {len(chart_rules)}")

    log("Extracting tables and charts from HTML...")
    contents = extract_content_from_input(
        input_path,
        chart_dir,
        report,
        chart_rules=chart_rules,
        allow_chart_fallback=allow_chart_fallback,
        cancel_check=cancel_check,
    )
    log(f"Extracted content blocks: {len(contents)}")
    _check_cancelled(cancel_check)
    if not contents:
        raise ValueError(f"No table, chart, or image content was found under HTML source folder: {input_path}")

    log("Resolving mappings...")
    registry = ContentRegistry(contents)
    resolved = resolve_mappings(rules, registry, report, cancel_check=cancel_check)
    if edb360_one_click:
        _apply_edb360_one_click_table_transforms(resolved)
        _apply_edb360_one_click_chart_groups(resolved, rules, contents)
    log(f"Resolved mappings: {len(resolved)} / {len(rules)}")
    _check_cancelled(cancel_check)

    if validate_only:
        log("Validation completed. No Word file was written.")
    else:
        log("Rendering Word report...")
        render_report(
            template_path,
            output_path,
            resolved,
            rules,
            report,
            log_callback=log_callback,
            text_mapping=text_mapping,
            max_table_rows=DEFAULT_MAX_TABLE_DATA_ROWS + 1,
            cleanup_unresolved_placeholders=edb360_one_click,
            cancel_check=cancel_check,
        )
        log(f"Done. Output saved to: {output_path}")

    _log_summary(report, log)
    return str(output_path)


def generate_edb360_report_to_file(
    html_input: str | Path,
    output_file: str | Path,
    metadata: dict[str, str] | None = None,
    mapping_file: str | Path = DEFAULT_MAPPING,
    master_template: str | Path = DEFAULT_EDB360_MASTER_TEMPLATE,
    report_language: str = "vi",
    chart_output_dir: str | Path | None = None,
    log_callback: LogCallback | None = None,
    cancel_check: Callable[[], bool] | None = None,
) -> str:
    """Generate an EDB360 Word report from the internal master template."""
    language = normalize_report_language(report_language)
    if language == "en" and Path(master_template) == DEFAULT_EDB360_MASTER_TEMPLATE:
        master_template = DEFAULT_EDB360_MASTER_TEMPLATE_EN
    template_path = _validate_word_file(master_template)
    input_path = _validate_html_input(html_input)
    detected_metadata = extract_edb360_metadata(input_path)
    merged_metadata = {**detected_metadata}
    for key, value in (metadata or {}).items():
        clean = str(value or "").strip()
        if clean:
            merged_metadata[key] = clean
    text_mapping = build_edb360_text_mapping(merged_metadata)
    text_mapping.update(build_edb360_assessment_mapping(input_path, language=language))
    log = _make_logger(log_callback)
    if detected_metadata:
        preview = ", ".join(f"{key}={value}" for key, value in sorted(detected_metadata.items()))
        log(f"Detected EDB360 metadata: {preview}")
    log(f"Report language: {language}")
    return generate_report_to_file(
        html_input=input_path,
        word_file=template_path,
        output_file=output_file,
        mapping_file=mapping_file,
        chart_output_dir=chart_output_dir,
        text_mapping=text_mapping,
        edb360_one_click=True,
        allow_chart_fallback=False,
        log_callback=log_callback,
        cancel_check=cancel_check,
    )


def _apply_edb360_one_click_table_transforms(
    resolved: dict[str, tuple[MappingRule, ExtractedContent]],
) -> None:
    if "<control_files>" in resolved:
        rule, content = resolved["<control_files>"]
        if isinstance(content, TableContent):
            resolved["<control_files>"] = (rule, _filter_parameter_rows(content, {"control_files"}))
    for placeholder, (_rule, content) in list(resolved.items()):
        if isinstance(content, TableContent) and _is_empty_edb360_table(content):
            del resolved[placeholder]


def _is_empty_edb360_table(content: TableContent) -> bool:
    if content.no_rows_selected and len(content.rows) <= 1:
        return True
    if not content.rows:
        return True
    if len(content.rows) == 1 and content.rows[0] == ["No rows selected"]:
        return True
    return len(content.rows) <= 1


def _apply_edb360_one_click_chart_groups(
    resolved: dict[str, tuple[MappingRule, ExtractedContent]],
    rules: list[MappingRule],
    contents: list[ExtractedContent],
) -> None:
    rules_by_placeholder = {rule.placeholder: rule for rule in rules}
    for placeholder, (generic_key, caption_template) in EDB360_CHART_GROUPS.items():
        items = _chart_group_items(contents, generic_key)
        if not items:
            continue
        captions = [caption_template.format(instance=_instance_number(item) or index + 1) for index, item in enumerate(items)]
        rule = rules_by_placeholder.get(placeholder) or MappingRule(
            placeholder=placeholder,
            source_key=generic_key,
            content_type="chart",
            width_inches=6.5,
        )
        resolved[placeholder] = (
            rule,
            MultiImageContent(
                source_path=items[0].source_path,
                images=items,
                captions=captions,
                title=generic_key,
                logical_key=generic_key,
                keys={generic_key},
            ),
        )


def _chart_group_items(contents: list[ExtractedContent], generic_key: str) -> list[ImageContent]:
    from utils.normalize import content_key_aliases

    matches: list[ImageContent] = []
    for content in contents:
        if not isinstance(content, ImageContent):
            continue
        aliases = {alias for key in content.keys | {content.logical_key, content.source_path.stem} for alias in content_key_aliases(key)}
        if generic_key in aliases:
            matches.append(content)
    matches = sorted(matches, key=lambda item: (_instance_number(item) or 999, item.source_path.name))
    if generic_key == "ash_top_sql_for_cluster":
        pie_matches = [item for item in matches if "pie_chart" in item.source_path.stem.lower()]
        return (pie_matches or matches)[:1]
    return matches


def _instance_number(content: ExtractedContent) -> int | None:
    text = " ".join([content.logical_key, content.title, content.source_path.stem])
    match = re.search(r"(?:for[_ ]instance|instance)[_ ]+(\d+)", text, re.IGNORECASE)
    return int(match.group(1)) if match else None


def _filter_parameter_rows(content: TableContent, parameter_names: set[str]) -> TableContent:
    if not content.rows:
        return content
    headers = content.rows[0]
    normalized_headers = [header.strip().upper().replace(" ", "_") for header in headers]
    try:
        name_index = normalized_headers.index("NAME")
    except ValueError:
        return content
    value_index = normalized_headers.index("VALUE") if "VALUE" in normalized_headers else None

    output_headers = ["PARAMETER", "VALUE"] if value_index is not None else headers
    output_rows = [output_headers]
    for row in content.rows[1:]:
        if name_index >= len(row) or row[name_index].strip().lower() not in parameter_names:
            continue
        if value_index is not None:
            values = [value.strip() for value in row[value_index].split(",") if value.strip()]
            output_rows.extend([[row[name_index], value] for value in values] or [[row[name_index], ""]])
        else:
            output_rows.append(row)

    transformed = copy(content)
    transformed.rows = output_rows
    transformed.no_rows_selected = len(output_rows) == 1
    return transformed


def _chart_rules_present_in_template(template_path: Path, rules):
    from docx import Document

    doc = Document(template_path)
    return [
        rule
        for rule in rules
        if rule.content_type == "chart" and document_contains_placeholder(doc, rule.placeholder)
    ]


def insert_placeholders_into_word(
    word_file: str | Path,
    mapping_file: str | Path = DEFAULT_MAPPING,
    log_callback: LogCallback | None = None,
    cancel_check: Callable[[], bool] | None = None,
) -> PlaceholderInsertReport:
    """Insert known report placeholders into a Word file in place."""
    template_path = _validate_word_file(word_file)
    mapping_path = Path(mapping_file)
    if not mapping_path.is_file():
        raise FileNotFoundError(f"Mapping file does not exist: {mapping_path}")

    log = _make_logger(log_callback)
    log("Loading placeholder mapping...")
    rules = load_mapping_rules(mapping_path)
    log(f"Loaded placeholders: {len(rules)}")
    log(f"Word file: {template_path}")
    _check_cancelled(cancel_check)

    report = insert_mapping_placeholders(template_path, rules, create_backup=True, cancel_check=cancel_check)
    if report.backup_path:
        log(f"Backup file: {report.backup_path}")
    log(f"Inserted placeholders: {len(report.inserted)}")
    for placeholder in report.inserted:
        log(f"  + {placeholder}")
    log(f"Already present: {len(report.already_present)}")
    if report.missing_anchors:
        log(f"Could not place placeholders: {len(report.missing_anchors)}")
        for placeholder in report.missing_anchors:
            log(f"  ! {placeholder}")
    else:
        log("All placeholders are present in the Word file.")
    return report


def run_sql_pipeline(
    input_root: str | Path,
    template_file: str | Path,
    output_root: str | Path | None = None,
    mapping_file: str | Path | None = DEFAULT_SQL_MAPPING,
    text_mapping: dict[str, str] | None = None,
    log_callback: LogCallback | None = None,
    cancel_check: Callable[[], bool] | None = None,
) -> list[str]:
    """Run SQLHealcheck CSV files -> merged Excel -> Word report."""
    input_path = _validate_sql_input_root(input_root)
    template_path = _validate_word_file(template_file)
    output_root_path = _validate_or_create_output_root(output_root or input_path)
    mapping_path = Path(mapping_file) if mapping_file else None
    log = _make_logger(log_callback)

    excel_file = output_root_path / "merged_healthcheck_info.xlsx"
    report_file = output_root_path / "final_healthcheck_report.docx"

    log("Running SQLHealcheck pipeline...")
    log(f"SQL root folder: {input_path}")
    log(f"Template: {template_path}")
    log(f"Selected output folder: {output_root_path}")
    log(f"Merged Excel output: {excel_file}")
    log(f"Word report output: {report_file}")
    _check_cancelled(cancel_check)

    merged_excel = merge_sql_root_healthcheck(input_path, excel_file, log_callback=log_callback, cancel_check=cancel_check)
    _check_cancelled(cancel_check)
    if not merged_excel:
        raise ValueError(f"No SQLHealcheck files were generated from: {input_path}")
    _check_cancelled(cancel_check)

    generated_report = render_excel_report(
        excel_path=merged_excel,
        template_path=template_path,
        output_path=report_file,
        mapping_path=mapping_path,
        text_mapping=text_mapping,
        log_callback=log_callback,
    )
    _check_cancelled(cancel_check)

    log("")
    log("SQLHealcheck completed.")
    log(f"Merged Excel file: {merged_excel}")
    log(f"Word report: {generated_report}")
    return [str(merged_excel), str(generated_report)]


def run_sql_one_click_pipeline(
    input_root: str | Path,
    output_root: str | Path,
    master_template: str | Path = DEFAULT_SQL_MASTER_TEMPLATE,
    mapping_file: str | Path | None = DEFAULT_SQL_MAPPING,
    metadata: dict[str, str] | None = None,
    report_language: str = "vi",
    log_callback: LogCallback | None = None,
    cancel_check: Callable[[], bool] | None = None,
) -> list[str]:
    """Run SQLHealthcheck reports from an extracted source root using the internal template."""
    input_path = Path(input_root)
    if not input_path.exists():
        raise FileNotFoundError(f"SQL source folder does not exist: {input_path}")
    if not input_path.is_dir():
        raise NotADirectoryError(f"SQL source must be a folder: {input_path}")

    language = normalize_report_language(report_language)
    if language == "en" and Path(master_template) == DEFAULT_SQL_MASTER_TEMPLATE:
        master_template = DEFAULT_SQL_MASTER_TEMPLATE_EN
    template_path = _validate_word_file(master_template)
    output_root_path = _validate_or_create_output_root(output_root)
    log = _make_logger(log_callback)

    sql_roots = _discover_sql_healthcheck_roots(input_path)
    if not sql_roots:
        raise ValueError(f"No SQLHealthcheck CSV folders were found under: {input_path}")

    log("Running SQLHealthcheck one-click pipeline...")
    log(f"Report language: {language}")
    log(f"Detected SQLHealthcheck folders: {len(sql_roots)}")
    generated_files: list[str] = []

    for index, sql_root in enumerate(sql_roots, start=1):
        _check_cancelled(cancel_check)
        report_name = _safe_report_stem(sql_root.name or f"sqlhealthcheck_{index}")
        report_output_root = output_root_path / report_name if len(sql_roots) > 1 else output_root_path
        log("")
        log(f"[{index}/{len(sql_roots)}] SQLHealthcheck folder: {sql_root}")
        report_files = run_sql_pipeline(
            input_root=sql_root,
            template_file=template_path,
            output_root=report_output_root,
            mapping_file=mapping_file,
            text_mapping=_sql_report_text_mapping(metadata, language=language),
            log_callback=log_callback,
            cancel_check=cancel_check,
        )
        generated_files.extend(report_files)

        generated_docx = report_output_root / "final_healthcheck_report.docx"
        named_docx = report_output_root / f"SGC_SQL_HEALTHCHECK_{report_name}.docx"
        if generated_docx.is_file() and generated_docx != named_docx:
            generated_docx.replace(named_docx)
            generated_files = [str(named_docx) if Path(path) == generated_docx else path for path in generated_files]

    log("")
    log("SQLHealthcheck one-click completed.")
    return generated_files


def _discover_sql_healthcheck_roots(input_path: Path) -> list[Path]:
    direct_csv_roots = sorted(
        {csv_file.parent for csv_file in input_path.rglob("*.csv") if csv_file.is_file()},
        key=lambda path: str(path).lower(),
    )
    if not direct_csv_roots:
        return []

    selected: list[Path] = []
    for candidate in direct_csv_roots:
        if any(parent in direct_csv_roots for parent in candidate.parents):
            continue
        selected.append(candidate)
    return selected


def _safe_report_stem(value: str) -> str:
    clean = re.sub(r"[^A-Za-z0-9_.-]+", "_", value).strip("._")
    return clean or "sqlhealthcheck"


def _sql_report_text_mapping(metadata: dict[str, str] | None, language: str = "vi") -> dict[str, str]:
    values = {
        "creator": "Trần Đinh Nhất Đăng",
        "approver": "Hồ Quốc Trí",
        "version": "1.0",
        "report_language": normalize_report_language(language),
        **(metadata or {}),
    }
    values["report_language"] = normalize_report_language(values.get("report_language", language))
    return {key: str(value).strip() for key, value in values.items() if str(value or "").strip()}


def normalize_report_language(value: str | None) -> str:
    text = str(value or "").strip().lower()
    if text in {"en", "eng", "english"}:
        return "en"
    return "vi"


def _check_cancelled(cancel_check: Callable[[], bool] | None) -> None:
    if cancel_check and cancel_check():
        raise OperationCancelled("Operation cancelled.")


def _make_logger(log_callback: LogCallback | None) -> LogCallback:
    def log(message: str) -> None:
        print(message)
        if log_callback:
            log_callback(message)

    return log


def _validate_html_root_folder(path_value: str | Path) -> Path:
    path = Path(path_value)
    if not path.exists():
        raise FileNotFoundError(f"HTML root folder does not exist: {path}")
    if not path.is_dir():
        raise NotADirectoryError(f"HTML source must be a root folder: {path}")
    return path


def _validate_html_input(path_value: str | Path) -> Path:
    path = Path(path_value)
    if not path.exists():
        raise FileNotFoundError(f"HTML input does not exist: {path}")
    if path.is_file() and path.suffix.lower() not in {".html", ".htm"}:
        raise ValueError(f"HTML input file must be .html or .htm: {path}")
    if not path.is_file() and not path.is_dir():
        raise ValueError(f"HTML input must be a file or folder: {path}")
    return path


def _validate_word_file(path_value: str | Path) -> Path:
    path = Path(path_value)
    if not path.exists():
        raise FileNotFoundError(f"Word file does not exist: {path}")
    if not path.is_file():
        raise ValueError(f"Word file path is not a file: {path}")
    if path.suffix.lower() != ".docx":
        raise ValueError(f"Word file must be a .docx file: {path}")
    return path


def _validate_sql_input_root(path_value: str | Path) -> Path:
    path = Path(path_value)
    if not path.exists():
        raise FileNotFoundError(f"SQL root folder does not exist: {path}")
    if not path.is_dir():
        raise NotADirectoryError(f"SQL input must be a root folder: {path}")
    if not _has_sql_healthcheck_input(path):
        raise ValueError(f"Selected SQL root folder must contain CSV files or DB subfolders with CSV files: {path}")
    return path


def _has_sql_healthcheck_input(path: Path) -> bool:
    if any(child.is_file() and child.suffix.lower() == ".csv" for child in path.iterdir()):
        return True
    return any(
        child.is_dir() and any(csv_file.is_file() for csv_file in child.glob("*.csv"))
        for child in path.iterdir()
    )


def _validate_or_create_output_root(path_value: str | Path) -> Path:
    path = Path(path_value)
    path.mkdir(parents=True, exist_ok=True)
    if not path.is_dir():
        raise NotADirectoryError(f"Output root is not a folder: {path}")
    return path


def _log_summary(report: GenerationReport, log: LogCallback) -> None:
    log("")
    log("Generation summary")
    log(f"Inserted: {len(report.inserted)}")
    for item in report.inserted:
        log(f"  + {item}")

    log(f"Missing content: {len(report.missing_content)}")
    for item in report.missing_content:
        log(f"  ! {item}")

    log(f"Missing placeholders: {len(report.missing_placeholders)}")
    for item in report.missing_placeholders:
        log(f"  ! {item}")

    log(f"Ambiguous mappings: {len(report.ambiguous)}")
    for item in report.ambiguous:
        log(f"  ! {item}")

    log(f"Warnings: {len(report.warnings)}")
    for item in report.warnings:
        log(f"  ! {item}")

    if report.skipped:
        log(f"Skipped: {len(report.skipped)}")
        for item in report.skipped:
            log(f"  ! {item}")
